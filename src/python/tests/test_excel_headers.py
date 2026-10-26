"""Strict header validation and output-equivalence tests using temporary workbooks."""

from contextlib import redirect_stderr, redirect_stdout
import importlib.util
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[3]
SPEC = importlib.util.spec_from_file_location(
    "header_generator", ROOT / "src/python/Excel_To_Html.py"
)
generator = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(generator)

COLUMNS = [
    "Property label", "Definition", "Property URI", "Range", "Cardinality",
    "Usage note", "Controlled vocabulary (if applicable)",
]


class ExcelHeaderTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.counter = 0

    def workbook(self, transform=lambda value: value, order=None, extra=False):
        self.counter += 1
        path = self.root / f"input-{self.counter}.xlsx"
        order = list(range(len(COLUMNS))) if order is None else order
        wb = Workbook()
        classes = wb.active
        classes.title = "classes"
        classes.append([transform("sheet_name"), "description"])
        classes.append(["Dataset", "Fixture class"])
        dataset = wb.create_sheet("Dataset")
        headers = [transform(COLUMNS[index]) for index in order]
        rows = [
            ["Example", "A definition\nwith a line break", "ex:example", "skos:Concept",
             "0..*", "Usage\nnote", "Example vocabulary"],
            [None, None, None, None, None, None,
             "The external vocabulary may be used for relevant terms."],
        ]
        if extra:
            headers += ["SHACL annotation", "notes", " notes ", None, 42]
        dataset.append(headers)
        for row in rows:
            values = [row[index] for index in order]
            if extra:
                values += ["Ignored source metadata"] * 5
            dataset.append(values)
        dataset.cell(2, order.index(6) + 1).hyperlink = "https://example.org/vocabulary"
        wb.save(path)
        wb.close()
        return path

    def change_header(self, path, sheet, column, value):
        wb = load_workbook(path)
        wb[sheet].cell(1, column).value = value
        wb.save(path)
        wb.close()

    def generate(self, path):
        output = self.root / f"output-{self.counter}"
        stdout, stderr = io.StringIO(), io.StringIO()
        with patch.object(generator, "OUTPUT_PATH", output):
            with redirect_stdout(stdout), redirect_stderr(stderr):
                try:
                    generator.main(["--excel", str(path)])
                    status = 0
                except SystemExit as exc:
                    status = exc.code
        files = {p.name: p.read_bytes() for p in output.glob("*.html")}
        return status, stdout.getvalue(), stderr.getvalue(), files

    def assert_rejected(self, path, worksheet, heading):
        status, stdout, stderr, files = self.generate(path)
        self.assertNotEqual(status, 0)
        self.assertIn(f"Error processing {worksheet}", stderr)
        self.assertIn(heading, stderr)
        self.assertIn("Available columns", stderr)
        self.assertNotIn("Generated ", stdout)
        self.assertEqual(files, {})

    def test_whitespace_normalisation_preserves_case_and_punctuation(self):
        for value in [
            "Property label", " Property label ", "Property  label",
            "Property\tlabel", "Property\r\nlabel", "Property\u00a0label",
        ]:
            with self.subTest(value=value):
                self.assertEqual(
                    generator.normalize_headers([value], "Dataset", ["Property label"]),
                    ["Property label"],
                )
        untouched = ["Property URl", "Property URI.", "sheet_name", None, 42]
        self.assertEqual(generator.normalize_headers(untouched, "Dataset", []), untouched)

    def test_correct_spelling_is_accepted_and_hyperlinks_are_preserved(self):
        status, _, stderr, files = self.generate(self.workbook())
        self.assertEqual(status, 0, stderr)
        html = files["properties-dataset.html"].decode()
        self.assertIn('<a href="https://example.org/vocabulary">', html)
        self.assertIn("The external vocabulary may be used for relevant terms.", html)

    def test_whitespace_variants_produce_identical_html(self):
        status, _, stderr, expected = self.generate(self.workbook())
        self.assertEqual(status, 0, stderr)
        transforms = [
            lambda value: f"  {value}  ",
            lambda value: value.replace(" ", "   "),
            lambda value: "\t" + value.replace(" ", "\t\n\u00a0") + "\r\n",
        ]
        for transform in transforms:
            with self.subTest(transform=transform):
                status, _, stderr, actual = self.generate(self.workbook(transform))
                self.assertEqual(status, 0, stderr)
                self.assertEqual(actual, expected)

    def test_missing_required_property_headings_are_rejected(self):
        for index, heading in enumerate(COLUMNS, 1):
            with self.subTest(heading=heading):
                path = self.workbook()
                self.change_header(path, "Dataset", index, "Unrelated heading")
                self.assert_rejected(path, "Dataset", heading)

    def test_misspelled_alternative_or_case_changed_headings_are_rejected(self):
        for index, value in [
            (7, "Controlled vocabluary (if applicable)"),
            (7, "Controlled vocabularies (if applicable)"),
            (7, "Controlled vocabulary"),
            (3, "Property URl"),
            (1, "property label"),
        ]:
            with self.subTest(value=value):
                path = self.workbook()
                self.change_header(path, "Dataset", index, value)
                self.assert_rejected(path, "Dataset", COLUMNS[index - 1])

    def test_duplicate_required_property_headings_are_rejected_before_pandas_renames_them(self):
        for heading in COLUMNS:
            for duplicate in [heading, f"\t {heading.replace(' ', '  ')} \u00a0"]:
                with self.subTest(heading=heading, duplicate=duplicate):
                    path = self.workbook()
                    self.change_header(path, "Dataset", 8, duplicate)
                    self.assert_rejected(path, "Dataset", heading)

    def test_missing_or_misspelled_classes_heading_is_rejected(self):
        for value in [None, "Sheet_name", "sheet name", "sheet_names"]:
            with self.subTest(value=value):
                path = self.workbook()
                self.change_header(path, "classes", 1, value)
                self.assert_rejected(path, "classes", "sheet_name")

    def test_duplicate_classes_heading_is_rejected(self):
        for duplicate in ["sheet_name", "\t sheet_name \u00a0"]:
            with self.subTest(duplicate=duplicate):
                path = self.workbook()
                self.change_header(path, "classes", 3, duplicate)
                self.assert_rejected(path, "classes", "sheet_name")

    def test_additional_unused_columns_are_accepted_but_not_output(self):
        status, _, stderr, expected = self.generate(self.workbook())
        self.assertEqual(status, 0, stderr)
        status, _, stderr, actual = self.generate(self.workbook(extra=True))
        self.assertEqual(status, 0, stderr)
        self.assertEqual(actual, expected)

    def test_reordered_columns_produce_identical_html(self):
        status, _, stderr, expected = self.generate(self.workbook())
        self.assertEqual(status, 0, stderr)
        status, _, stderr, actual = self.generate(
            self.workbook(order=[6, 4, 2, 0, 5, 3, 1])
        )
        self.assertEqual(status, 0, stderr)
        self.assertEqual(actual, expected)

    def test_openpyxl_uses_same_normalisation_for_hyperlink_lookup(self):
        canonical = self.workbook()
        formatted = self.workbook(lambda value: " \t" + value.replace(" ", "\n") + "\u00a0")
        self.assertEqual(
            generator.extract_urls_per_property("Dataset", formatted),
            generator.extract_urls_per_property("Dataset", canonical),
        )

    def test_openpyxl_rejects_duplicate_or_invalid_required_headings(self):
        for heading in ["Property URI", "Controlled vocabulary (if applicable)"]:
            path = self.workbook()
            self.change_header(path, "Dataset", 8, heading)
            with self.subTest(heading=heading):
                with self.assertRaisesRegex(ValueError, "occurs 2 times"):
                    generator.extract_urls_per_property("Dataset", path)
        path = self.workbook()
        self.change_header(path, "Dataset", 7, "Controlled vocabluary (if applicable)")
        with self.assertRaisesRegex(ValueError, "was not found"):
            generator.extract_urls_per_property("Dataset", path)


if __name__ == "__main__":
    unittest.main()
