"""Check source values and links in freshly generated property tables.

Run with: python -m unittest discover -s src/python -p 'test_*.py'
"""

from html.parser import HTMLParser
from pathlib import Path
import importlib.util
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "src/python/Excel_To_Html.py"
spec = importlib.util.spec_from_file_location("excel_generator", SCRIPT)
generator = importlib.util.module_from_spec(spec)
spec.loader.exec_module(generator)


def text(value):
    return re.sub(r"\s+", " ", str(value)).strip()


class TableReader(HTMLParser):
    def __init__(self, source):
        super().__init__(convert_charrefs=True)
        self.rows = []
        self.cell = None
        self.row = None
        self.feed(source)

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag == "td":
            self.cell = {"text": "", "links": []}
        elif self.cell is not None and tag == "a":
            self.cell["links"].append(dict(attrs).get("href"))
        elif self.cell is not None and tag == "br":
            self.cell["text"] += " "

    def handle_data(self, data):
        if self.cell is not None:
            self.cell["text"] += data

    def handle_endtag(self, tag):
        if tag == "td" and self.cell is not None:
            self.row.append(self.cell)
            self.cell = None
        elif tag == "tr" and self.row:
            self.rows.append(self.row)
            self.row = None


class GeneratedTables(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.build = Path(cls.temp.name)
        script_dir = cls.build / "src/python"
        script_dir.mkdir(parents=True)
        for name in ["Excel_To_Html.py", "links.json"]:
            shutil.copy2(SCRIPT.parent / name, script_dir / name)
        workbook = (SCRIPT.parent / generator.EXCEL_FILE_PATH).resolve()
        selected = (script_dir / generator.EXCEL_FILE_PATH).resolve()
        selected.parent.mkdir(parents=True)
        shutil.copy2(workbook, selected)
        cls.book = load_workbook(workbook)
        cls.sheets = [row[0].value for row in cls.book["classes"].iter_rows(min_row=2)]
        subprocess.run([sys.executable, "Excel_To_Html.py"], cwd=script_dir,
                       check=True, capture_output=True, text=True)
        cls.tables = {
            sheet: TableReader((cls.build / "src/property" /
                                f"properties-{sheet.lower()}.html").read_text()).rows
            for sheet in cls.sheets
        }

    @classmethod
    def tearDownClass(cls):
        cls.book.close()
        cls.temp.cleanup()

    def source_rows(self, sheet):
        ws = self.book[sheet]
        headers = {cell.value: cell.column for cell in ws[1] if cell.value}
        rows = [row for row in ws.iter_rows(min_row=2)
                if row[headers["Property URI"] - 1].value]
        return headers, rows

    def test_expected_tables_and_property_rows(self):
        actual = {p.name for p in (self.build / "src/property").glob("*.html")}
        self.assertEqual(actual, {f"properties-{s.lower()}.html" for s in self.sheets})
        for sheet in self.sheets:
            with self.subTest(sheet=sheet):
                _, rows = self.source_rows(sheet)
                self.assertEqual(len(self.tables[sheet]), len(rows))

    def test_property_values_and_usage_notes_survive_generation(self):
        columns = ["Property label", "Definition", "Property URI", "Range", "Cardinality"]
        for sheet in self.sheets:
            headers, rows = self.source_rows(sheet)
            for source, generated in zip(rows, self.tables[sheet]):
                with self.subTest(sheet=sheet, row=source[0].row):
                    self.assertEqual(len(generated), 6)
                    for index, column in enumerate(columns):
                        value = source[headers[column] - 1].value
                        expected = "NaN" if value is None else text(value)
                        self.assertEqual(text(generated[index]["text"]), expected)
                    usage = source[headers["Usage note"] - 1].value
                    expected = "N.A." if not usage else text(usage.replace("\\n", " "))
                    self.assertIn(expected, text(generated[5]["text"]))

    def test_source_cell_hyperlinks_survive_generation(self):
        columns = ["Property label", "Definition", "Property URI", "Range", "Cardinality", "Usage note"]
        for sheet in self.sheets:
            headers, rows = self.source_rows(sheet)
            for source, generated in zip(rows, self.tables[sheet]):
                for index, column in enumerate(columns):
                    cell = source[headers[column] - 1]
                    if cell.hyperlink:
                        with self.subTest(sheet=sheet, cell=cell.coordinate):
                            self.assertIn(cell.hyperlink.target, generated[index]["links"])

    def test_checked_in_hyperlinks_stay_with_their_source_properties(self):
        columns = ["Property label", "Definition", "Property URI", "Range", "Cardinality", "Usage note"]
        for sheet in self.sheets:
            path = ROOT / "src/property" / f"properties-{sheet.lower()}.html"
            checked = TableReader(path.read_text()).rows
            headers, rows = self.source_rows(sheet)
            for source, generated in zip(rows, checked):
                for index, column in enumerate(columns):
                    cell = source[headers[column] - 1]
                    if cell.hyperlink:
                        with self.subTest(sheet=sheet, cell=cell.coordinate):
                            self.assertIn(cell.hyperlink.target, generated[index]["links"])

    def test_column_property_reference_remains_absolute(self):
        path = ROOT / "src/property/properties-column.html"
        row = next(r for r in TableReader(path.read_text()).rows
                   if text(r[2]["text"]) == "csvw:propertyUrl")
        self.assertIn("https://www.w3.org/TR/tabular-metadata/#propertyurl", row[5]["links"])

    def test_availability_does_not_gain_an_optional_vocabulary_statement(self):
        row = next(r for r in self.tables["Distribution"]
                   if text(r[2]["text"]) == "dcatap:availability")
        usage = text(row[5]["text"])
        self.assertIn("MUST be used", usage)
        self.assertNotIn("vocabulary is optional", usage)

    def test_cell_links_preserve_special_characters(self):
        label = '<urn:example> & café "quoted"'
        target = 'https://example.org/?a=1&b="two"'
        parsed = TableReader('<table><tr><td>' +
                             generator.make_cell_link(label, target) +
                             '</td></tr></table>').rows[0][0]
        self.assertEqual(parsed["text"], label)
        self.assertEqual(parsed["links"], [target])


class VersionReferences(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        for path in ["index.bs", "src/chapter/4-introduction.md",
                     "src/python/Excel_To_Html.py", "src/python/check_version_consistency.py",
                     "src/excel/HealthRI_v2.1.0.xlsx", "src/images/HRI_metadata_V2.1.png"]:
            dest = self.root / path
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ROOT / path, dest)

    def run_check(self):
        return subprocess.run([sys.executable, "src/python/check_version_consistency.py"],
                              cwd=self.root, capture_output=True, text=True)

    def test_current_source_references_pass(self):
        result = self.run_check()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_mismatched_diagram_version_is_rejected(self):
        intro = self.root / "src/chapter/4-introduction.md"
        intro.write_text(intro.read_text().replace("HRI_metadata_V2.1.png", "HRI_metadata_V2.0.png"))
        shutil.copy2(self.root / "src/images/HRI_metadata_V2.1.png",
                     self.root / "src/images/HRI_metadata_V2.0.png")
        result = self.run_check()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Major/minor mismatch", result.stderr)

    def test_missing_diagram_is_rejected(self):
        (self.root / "src/images/HRI_metadata_V2.1.png").unlink()
        result = self.run_check()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("referenced UML image does not exist", result.stderr)


if __name__ == "__main__":
    unittest.main()
